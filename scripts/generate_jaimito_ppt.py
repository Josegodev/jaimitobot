from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT_PATH = Path("Jaimito_infografia_canva.pptx")

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BLUE_LIGHT = RGBColor(217, 238, 255)
BLUE = RGBColor(74, 144, 226)
GREEN = RGBColor(54, 179, 126)
WHITE = RGBColor(255, 255, 255)
GRAY_DARK = RGBColor(47, 58, 69)
GRAY_SOFT = RGBColor(238, 243, 247)


def set_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(
    slide,
    text: str,
    left,
    top,
    width,
    height,
    *,
    font_size: int = 20,
    bold: bool = False,
    color: RGBColor = GRAY_DARK,
    align=PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_panel(slide, left, top, width, height, *, fill=WHITE, line=BLUE) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line


def add_bullet_list(
    slide,
    items: list[str],
    left,
    top,
    width,
    height,
    *,
    font_size: int = 18,
    color: RGBColor = GRAY_DARK,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        run = paragraph.runs[0]
        run.font.size = Pt(font_size)
        run.font.color.rgb = color


def add_flow_box(slide, title: str, body: str, left, top, width, height, *, fill) -> None:
    add_panel(slide, left, top, width, height, fill=fill, line=BLUE)
    add_textbox(
        slide,
        title,
        left + Inches(0.15),
        top + Inches(0.1),
        width - Inches(0.3),
        Inches(0.35),
        font_size=18,
        bold=True,
        color=GRAY_DARK,
    )
    add_textbox(
        slide,
        body,
        left + Inches(0.15),
        top + Inches(0.45),
        width - Inches(0.3),
        height - Inches(0.55),
        font_size=12,
        color=GRAY_DARK,
    )


def add_arrow_text(slide, text: str, left, top, width, height) -> None:
    add_textbox(
        slide,
        text,
        left,
        top,
        width,
        height,
        font_size=20,
        bold=True,
        color=GREEN,
        align=PP_ALIGN.CENTER,
    )


def build_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)

    add_panel(slide, Inches(0.6), Inches(0.7), Inches(12.1), Inches(5.8), fill=BLUE_LIGHT)
    add_textbox(
        slide,
        "Como funciona Jaimito?",
        Inches(1.0),
        Inches(1.3),
        Inches(8.0),
        Inches(0.8),
        font_size=28,
        bold=True,
        color=BLUE,
    )
    add_textbox(
        slide,
        "El asistente del curso de IA aplicada que ayuda a preparar la web del proyecto final",
        Inches(1.0),
        Inches(2.0),
        Inches(9.5),
        Inches(0.7),
        font_size=18,
        color=GRAY_DARK,
    )
    add_panel(slide, Inches(9.8), Inches(1.4), Inches(2.0), Inches(2.0), fill=WHITE, line=GREEN)
    add_textbox(
        slide,
        "JAIMITO",
        Inches(10.05),
        Inches(2.15),
        Inches(1.5),
        Inches(0.5),
        font_size=20,
        bold=True,
        color=GREEN,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        "Bot de Telegram para ayudar al alumnado paso a paso.",
        Inches(1.0),
        Inches(3.1),
        Inches(7.5),
        Inches(0.8),
        font_size=20,
        color=GRAY_DARK,
    )
    add_textbox(
        slide,
        "Mensaje clave: Jaimito no sustituye al profesor. Ayuda a avanzar con claridad.",
        Inches(1.0),
        Inches(4.3),
        Inches(10.5),
        Inches(0.6),
        font_size=16,
        bold=True,
        color=GREEN,
    )


def build_flow_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)

    add_textbox(
        slide,
        "Flujo principal",
        Inches(0.6),
        Inches(0.3),
        Inches(3.0),
        Inches(0.5),
        font_size=24,
        bold=True,
        color=BLUE,
    )

    box_w = Inches(1.9)
    box_h = Inches(1.55)
    top = Inches(2.1)
    xs = [Inches(0.35), Inches(2.55), Inches(4.75), Inches(6.95), Inches(9.15), Inches(11.35)]
    blocks = [
        ("Alumno en Telegram", "El alumno escribe su duda y pide ayuda con comandos sencillos.", BLUE_LIGHT),
        ("Jaimito Bot", "Lee el mensaje y decide si responde directo o si consulta IA.", WHITE),
        ("Railway 24/7", "Servidor en la nube que mantiene el bot encendido.", BLUE_LIGHT),
        ("Docker", "Caja con todo lo necesario para que funcione igual siempre.", WHITE),
        ("OpenAI", "Genera respuestas cuando hace falta una ayuda mas elaborada.", BLUE_LIGHT),
        ("Respuesta", "El alumno recibe orientacion practica en Telegram.", WHITE),
    ]

    for index, (title, body, fill) in enumerate(blocks):
        add_flow_box(slide, title, body, xs[index], top, box_w, box_h, fill=fill)

    for left in [Inches(2.0), Inches(4.2), Inches(6.4), Inches(8.6), Inches(10.8)]:
        add_arrow_text(slide, ">", left, Inches(2.65), Inches(0.35), Inches(0.3))

    add_panel(slide, Inches(0.8), Inches(5.2), Inches(11.8), Inches(1.4), fill=GRAY_SOFT, line=GREEN)
    add_textbox(
        slide,
        "Ejemplos de uso: /pregunta Como empiezo mi web?   /web Quiero hacer una idea sobre una tienda",
        Inches(1.1),
        Inches(5.65),
        Inches(11.2),
        Inches(0.5),
        font_size=16,
        color=GRAY_DARK,
    )


def build_files_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)

    add_textbox(
        slide,
        "Por dentro de Jaimito",
        Inches(0.6),
        Inches(0.3),
        Inches(4.5),
        Inches(0.5),
        font_size=24,
        bold=True,
        color=BLUE,
    )
    add_textbox(
        slide,
        "Cuatro archivos clave explicados con lenguaje sencillo",
        Inches(0.6),
        Inches(0.8),
        Inches(6.2),
        Inches(0.4),
        font_size=14,
        color=GRAY_DARK,
    )

    positions = [
        (Inches(0.6), Inches(1.5)),
        (Inches(6.8), Inches(1.5)),
        (Inches(0.6), Inches(4.15)),
        (Inches(6.8), Inches(4.15)),
    ]
    cards = [
        (
            "config.py",
            "Prepara la configuracion del bot. Lee claves y decide el modelo de IA.",
            "openai_model=os.getenv(\"OPENAI_MODEL\", DEFAULT_MODEL)",
        ),
        (
            "jaimito_bot.py",
            "Es el archivo principal. Recibe comandos, llama a OpenAI y envia respuestas.",
            "application.add_handler(CommandHandler(\"web\", web))",
        ),
        (
            "logging_utils.py",
            "Guarda registros internos para revisar errores y entender que paso.",
            "def new_trace_id() -> str: return uuid4().hex",
        ),
        (
            "prompts.py",
            "Define como debe hablar Jaimito y que tipo de ayuda debe dar.",
            "SYSTEM_PROMPT = \"Eres Jaimito...\"",
        ),
    ]

    for (left, top), (title, body, code) in zip(positions, cards):
        add_panel(slide, left, top, Inches(5.7), Inches(2.25), fill=GRAY_SOFT, line=BLUE)
        add_textbox(slide, title, left + Inches(0.2), top + Inches(0.15), Inches(2.4), Inches(0.35), font_size=18, bold=True, color=BLUE)
        add_textbox(slide, body, left + Inches(0.2), top + Inches(0.55), Inches(5.1), Inches(0.7), font_size=14, color=GRAY_DARK)
        add_panel(slide, left + Inches(0.2), top + Inches(1.35), Inches(5.1), Inches(0.55), fill=WHITE, line=GREEN)
        add_textbox(slide, code, left + Inches(0.3), top + Inches(1.48), Inches(4.9), Inches(0.3), font_size=12, color=GRAY_DARK)


def build_capabilities_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)

    add_textbox(
        slide,
        "Que puede hacer Jaimito",
        Inches(0.6),
        Inches(0.3),
        Inches(5.0),
        Inches(0.5),
        font_size=24,
        bold=True,
        color=BLUE,
    )

    items = [
        "Ayudar a elegir una idea de proyecto",
        "Explicar las secciones de la web",
        "Redactar textos claros",
        "Dar una plantilla de web",
        "Explicar la rubrica",
        "Resolver dudas durante el curso",
    ]

    card_positions = [
        (Inches(0.8), Inches(1.3)),
        (Inches(4.55), Inches(1.3)),
        (Inches(8.3), Inches(1.3)),
        (Inches(0.8), Inches(3.8)),
        (Inches(4.55), Inches(3.8)),
        (Inches(8.3), Inches(3.8)),
    ]

    for idx, ((left, top), text) in enumerate(zip(card_positions, items), start=1):
        add_panel(slide, left, top, Inches(3.1), Inches(1.8), fill=BLUE_LIGHT if idx % 2 else WHITE, line=BLUE)
        add_textbox(
            slide,
            f"{idx}.",
            left + Inches(0.18),
            top + Inches(0.18),
            Inches(0.35),
            Inches(0.35),
            font_size=18,
            bold=True,
            color=GREEN,
        )
        add_textbox(
            slide,
            text,
            left + Inches(0.55),
            top + Inches(0.18),
            Inches(2.3),
            Inches(1.2),
            font_size=16,
            color=GRAY_DARK,
        )


def build_control_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)

    add_textbox(
        slide,
        "Como se controla",
        Inches(0.6),
        Inches(0.3),
        Inches(4.0),
        Inches(0.5),
        font_size=24,
        bold=True,
        color=BLUE,
    )

    add_panel(slide, Inches(0.7), Inches(1.2), Inches(6.0), Inches(4.8), fill=GRAY_SOFT, line=BLUE)
    add_bullet_list(
        slide,
        [
            "Solo responde a comandos o menciones.",
            "No responde a todos los mensajes del grupo.",
            "Las claves se guardan como variables privadas.",
            "Los errores se revisan en registros internos.",
            "No se muestra informacion tecnica al alumnado.",
        ],
        Inches(1.0),
        Inches(1.7),
        Inches(5.2),
        Inches(3.7),
        font_size=18,
    )

    add_panel(slide, Inches(7.1), Inches(1.2), Inches(5.4), Inches(4.8), fill=BLUE_LIGHT, line=GREEN)
    add_textbox(
        slide,
        "Idea sencilla para explicar esto en clase",
        Inches(7.4),
        Inches(1.55),
        Inches(4.7),
        Inches(0.5),
        font_size=18,
        bold=True,
        color=GREEN,
    )
    add_textbox(
        slide,
        "Jaimito ayuda, pero no habla cuando no toca. Ademas, protege la parte interna para que el alumnado solo vea una respuesta clara y util.",
        Inches(7.4),
        Inches(2.15),
        Inches(4.5),
        Inches(1.8),
        font_size=18,
        color=GRAY_DARK,
    )


def build_close_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BLUE_LIGHT)

    add_panel(slide, Inches(1.0), Inches(1.25), Inches(11.2), Inches(4.7), fill=WHITE, line=GREEN)
    add_textbox(
        slide,
        "Mensaje final",
        Inches(1.4),
        Inches(1.8),
        Inches(2.5),
        Inches(0.5),
        font_size=20,
        bold=True,
        color=GREEN,
    )
    add_textbox(
        slide,
        "Jaimito no sustituye al profesor: ayuda al alumnado a avanzar paso a paso con su proyecto.",
        Inches(1.4),
        Inches(2.55),
        Inches(10.0),
        Inches(1.0),
        font_size=28,
        bold=True,
        color=BLUE,
        align=PP_ALIGN.CENTER,
    )
    add_textbox(
        slide,
        "Presentacion editable para seguir ajustando en PowerPoint o subirla despues a Canva.",
        Inches(2.2),
        Inches(4.5),
        Inches(8.5),
        Inches(0.5),
        font_size=16,
        color=GRAY_DARK,
        align=PP_ALIGN.CENTER,
    )


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    build_title_slide(prs)
    build_flow_slide(prs)
    build_files_slide(prs)
    build_capabilities_slide(prs)
    build_control_slide(prs)
    build_close_slide(prs)

    prs.save(OUTPUT_PATH)
    print(f"Created {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
